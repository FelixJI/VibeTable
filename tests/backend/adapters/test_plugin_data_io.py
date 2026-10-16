from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook

from backend.adapters.pocketbase.client import PocketBaseClient
from backend.adapters.pocketbase.data_io import ProductDataIoRuntime
from backend.application.host_files import HostFiles
from backend.application.task_runtime import HostExecutionRuntime
from backend.application.task_service import TaskService
from backend.contracts.task import StartTaskExecutionParams
from tests.backend.host_files_fixture import LocalFilePeer


class ScopedClient(PocketBaseClient):
    def __init__(self):
        self.calls = []
        self.revoked = False

    async def plugin_surface(self, principal, method, request, **kwargs):
        self.calls.append((principal, method, request, kwargs))
        if self.revoked:
            raise PermissionError("installation revoked")
        assert principal == {
            "installationId": "host-owned",
            "surfaceToken": "surface",
            "installationRevision": 3,
        }
        if method == "data.ioProfile":
            return {
                "collection": "orders",
                "schema_revision": "schema_1",
                "fields": ["id", "title"],
                # The scoped export header must stay the authorized plugin field ID; the
                # display name carried by the profile never becomes a column identity.
                "field_schemas": {"title": {"dataType": "text", "displayName": "订单标题"}},
                "create_fields": ["title"],
                "update_fields": ["title"],
                "archive_field": None,
                "date_updated_field": None,
            }
        if method == "import.preview":
            return {"rows": [{"values": row, "diagnostics": []} for row in request["rows"]]}
        if method == "import.mint":
            assert request["rows"][0]["values"] == {"title": "synthetic"}
            return {"token": "imp1.host-owned", "expiresAt": 9999999999}
        assert method == "data.exportQuery"
        assert request["fields"] == ["title"]
        return {
            "items": [{"title": f"row-{index}", "id": "must-not-export"} for index in range(100)],
            "nextCursor": None,
            "schemaRevision": "schema_1",
            "dataRevision": 4,
            "clockPeriod": "2026-10-15T09:00Z",
            "filteredRows": 100,
            "totalRows": 100,
            "dependencies": {},
            "complete": True,
        }


def scope():
    return {
        "principal": {
            "installationId": "host-owned",
            "surfaceToken": "surface",
            "installationRevision": 3,
        },
        "runId": "host-run",
        "fields": ["title"],
        "query": {},
    }


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
async def test_scoped_export_reuses_real_task_and_formatter_without_hidden_id(tmp_path, fmt):
    peer = LocalFilePeer()
    path = tmp_path / f"output.{fmt}"
    grant = peer.grants.issue(purpose="export_target", direction="write", path=str(path))
    # The fixture store returns a descriptor in the same shape as native grants.
    grant_id = grant.grant_id
    peer.runs[grant_id] = "host-run"
    original = peer.call
    calls = []

    async def call(action, params):
        calls.append((action, params))
        return await original(action, params)

    reports = []

    async def sink(value):
        reports.append(value)

    runtime = HostExecutionRuntime(notification_sink=sink)
    tasks = TaskService(runtime, HostFiles(call))
    client = ScopedClient()
    ProductDataIoRuntime(client=client, task_service=tasks).register_tasks()
    await tasks.start_execution(
        StartTaskExecutionParams(
            task_id="owned-task",
            kind="data.export",
            params={
                "grantId": grant_id,
                "collection": "orders",
                "format": fmt,
                "pluginScope": scope(),
            },
        )
    )
    import asyncio

    for _ in range(100):
        if reports and reports[-1].state in {"succeeded", "failed"}:
            break
        await asyncio.sleep(0.01)
    assert reports[-1].state == "succeeded", reports[-1].error
    assert reports[-1].result.rows_written == 100
    assert sum(method == "data.exportQuery" for _, method, _, _ in client.calls) == 1
    assert next(params for action, params in calls if action == "openWrite")["runId"] == "host-run"
    if fmt == "csv":
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
    else:
        book = load_workbook(path, read_only=True)
        try:
            rows = [list(row) for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
    assert rows[0] == ["title"]
    assert len(rows) == 101
    assert "must-not-export" not in str(rows)


async def test_result_shape_drift_aborts_before_commit_and_reports_not_committed(tmp_path):
    """A result-contract drift must abort the export before the native commit.

    CollectionProfile accepts a schema revision longer than the frozen
    ExportResult limit. Constructing the typed result before the commit turns
    that drift into a definite pre-commit failure — finishWrite(commit=False),
    target untouched, positive not-committed receipt — instead of a committed
    file mislabeled as not committed.
    """
    import asyncio

    peer = LocalFilePeer()
    path = tmp_path / "drift.csv"
    grant = peer.grants.issue(purpose="export_target", direction="write", path=str(path))
    peer.runs[grant.grant_id] = "host-run"
    original = peer.call
    calls = []

    async def call(action, params):
        calls.append((action, params))
        return await original(action, params)

    reports = []

    async def sink(value):
        reports.append(value)

    class LongRevisionClient(ScopedClient):
        async def plugin_surface(self, principal, method, request, **kwargs):
            result = await super().plugin_surface(principal, method, request, **kwargs)
            if method == "data.ioProfile":
                result["schema_revision"] = "r" * 129
            return result

    runtime = HostExecutionRuntime(notification_sink=sink)
    tasks = TaskService(runtime, HostFiles(call))
    ProductDataIoRuntime(client=LongRevisionClient(), task_service=tasks).register_tasks()
    await tasks.start_execution(
        StartTaskExecutionParams(
            task_id="owned-task",
            kind="data.export",
            params={
                "grantId": grant.grant_id,
                "collection": "orders",
                "format": "csv",
                "pluginScope": scope(),
            },
        )
    )
    for _ in range(200):
        if reports and reports[-1].state in {"succeeded", "failed"}:
            break
        await asyncio.sleep(0.01)
    terminal = reports[-1]
    assert terminal.state == "failed", terminal.error
    assert terminal.commit_receipt == "not_committed"
    assert not path.exists(), "the drift must abort before the target file is replaced"
    assert [
        params for action, params in calls if action == "finishWrite" and params["commit"] is True
    ] == []
    assert any(action == "finishWrite" and params["commit"] is False for action, params in calls)


async def test_scoped_import_uses_preview_and_go_plan_owner_without_platform_apply(tmp_path):
    from backend.adapters.pocketbase.plugin_data_io import preview_plugin_import

    peer = LocalFilePeer()
    source = tmp_path / "source.csv"
    source.write_text("title\nsynthetic\n", encoding="utf-8")
    grant = peer.grants.issue(purpose="import_source", direction="read", path=str(source))
    peer.runs[grant.grant_id] = "host-run"
    client = ScopedClient()

    async def progress(*_args):
        pass

    plan = await preview_plugin_import(
        client,
        peer.files,
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "pluginScope": scope(),
        },
        progress,
        lambda: False,
    )
    assert plan.rows[0].values == {"title": "synthetic"}
    assert [method for _, method, _, _ in client.calls] == [
        "data.ioProfile",
        "import.preview",
        "import.mint",
    ]


async def test_scoped_upsert_preview_matches_by_key_and_surfaces_missing_required_rows(tmp_path):
    """The full plugin chain must deliver the approved upsert key to the
    authoritative preview so unmatched rows normalize as inserts; otherwise a
    missing required field silently bypasses the checkable errorRows."""
    from backend.adapters.pocketbase.plugin_data_io import preview_plugin_import

    class UpsertClient(ScopedClient):
        """Contract-faithful import.preview peer: insert mode checks required
        fields, update mode does not, and upsert rows only update when the
        approved key matches an existing record."""

        def __init__(self):
            super().__init__()
            self.existing_keys = {"existing"}
            self.required_fields = ("amount",)

        async def plugin_surface(self, principal, method, request, **kwargs):
            if method in {"data.ioProfile", "import.preview", "import.mint"}:
                self.calls.append((principal, method, request, kwargs))
            if method == "data.ioProfile":
                return {
                    "collection": "orders",
                    "schema_revision": "schema_1",
                    "fields": ["id", "key", "amount", "title"],
                    "field_schemas": {
                        "key": {"dataType": "text", "displayName": "key"},
                        "amount": {"dataType": "number", "displayName": "amount"},
                        "title": {"dataType": "text", "displayName": "title"},
                    },
                    "create_fields": ["key", "amount", "title"],
                    "update_fields": ["key", "amount", "title"],
                    "archive_field": None,
                    "date_updated_field": None,
                }
            if method == "import.preview":
                upsert_key = request.get("upsertKey")
                rows_out = []
                for row in request["rows"]:
                    mode = "update" if request["mode"] == "upsert" else "insert"
                    if request["mode"] == "upsert" and upsert_key:
                        matched = row.get(upsert_key) in self.existing_keys
                        mode = "update" if matched else "insert"
                    diagnostics = []
                    if mode == "insert":
                        for field in self.required_fields:
                            if row.get(field) is None or row.get(field) == "":
                                diagnostics.append(
                                    {
                                        "field": field,
                                        "code": "field.value.required",
                                        "message": "required field is not supplied",
                                    }
                                )
                    rows_out.append({"values": row, "diagnostics": diagnostics})
                return {"rows": rows_out}
            if method == "import.mint":
                return {"token": "imp1.host-owned", "expiresAt": 9999999999}
            return await super().plugin_surface(principal, method, request, **kwargs)

    peer = LocalFilePeer()
    source = tmp_path / "upsert.csv"
    source.write_text("key,amount,title\nexisting,10,keep\nfresh,,new\n", encoding="utf-8")
    grant = peer.grants.issue(purpose="import_source", direction="read", path=str(source))
    peer.runs[grant.grant_id] = "host-run"
    client = UpsertClient()

    async def progress(*_args):
        pass

    plan = await preview_plugin_import(
        client,
        peer.files,
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "mode": "upsert",
            "upsertKey": "key",
            "pluginScope": scope(),
        },
        progress,
        lambda: False,
    )
    # The matched row keeps partial update semantics; the new insert row
    # surfaces its missing required field as a checkable diagnostic. These
    # semantic checks come first: without the key the whole upsert batch is
    # silently normalized as updates and the missing required field never
    # enters the checkable errorRows.
    assert plan.rows[0].diagnostics == []
    assert any(diagnostic.code == "field.value.required" for diagnostic in plan.rows[1].diagnostics)
    assert plan.summary.error_rows == 1
    preview_call = None
    for _, method, request, _ in client.calls:
        if method == "import.preview":
            preview_call = request
            break
    assert preview_call is not None
    # The approved key reaches the authoritative preview on the public wire.
    assert preview_call["upsertKey"] == "key"


async def test_scoped_create_only_preview_keeps_insert_required_checks_and_null_key(tmp_path):
    from backend.adapters.pocketbase.plugin_data_io import preview_plugin_import

    class CreateOnlyClient(ScopedClient):
        async def plugin_surface(self, principal, method, request, **kwargs):
            if method == "import.preview":
                assert request["mode"] == "create_only"
                # Ordinary inserts always run the required-field check; the
                # key is explicitly absent instead of silently defaulted.
                assert "upsertKey" in request
                assert request["upsertKey"] is None
                diagnostics = (
                    [
                        {
                            "field": "title",
                            "code": "field.value.required",
                            "message": "required field is not supplied",
                        }
                    ]
                    if any(row.get("title") is None for row in request["rows"])
                    else []
                )
                return {
                    "rows": [
                        {"values": row, "diagnostics": list(diagnostics)} for row in request["rows"]
                    ]
                }
            if method == "import.mint":
                return {"token": "imp1.host-owned", "expiresAt": 9999999999}
            return await super().plugin_surface(principal, method, request, **kwargs)

    peer = LocalFilePeer()
    source = tmp_path / "create.csv"
    source.write_text("title\n\n", encoding="utf-8")
    grant = peer.grants.issue(purpose="import_source", direction="read", path=str(source))
    peer.runs[grant.grant_id] = "host-run"

    async def progress(*_args):
        pass

    plan = await preview_plugin_import(
        CreateOnlyClient(),
        peer.files,
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "mode": "create_only",
            "pluginScope": scope(),
        },
        progress,
        lambda: False,
    )
    assert plan.summary.error_rows == 1


async def test_scoped_upsert_maps_operation_specific_columns_and_denies_wrong_operation_rows(
    tmp_path,
):
    """The full plugin chain routes columns granted by either import operation:
    an update-only column reaches the authoritative preview, a matched row
    keeps it clean, and an unmatched insert row gets the Go per-operation
    field-grant diagnostic instead of silently dropping the column."""
    from backend.adapters.pocketbase.plugin_data_io import preview_plugin_import

    class HeterogeneousGrantClient(ScopedClient):
        """Contract-faithful import.preview peer mirroring the fixed Go surface:
        rows are matched by the approved key, then authorized against their
        actual operation grant exactly like the authoritative preview/mint."""

        def __init__(self):
            super().__init__()
            self.existing_keys = {"existing"}
            self.create_fields = {"key", "title"}
            self.update_fields = {"key", "note"}

        async def plugin_surface(self, principal, method, request, **kwargs):
            if method in {"data.ioProfile", "import.preview", "import.mint"}:
                self.calls.append((principal, method, request, kwargs))
            if method == "data.ioProfile":
                return {
                    "collection": "orders",
                    "schema_revision": "schema_1",
                    "fields": ["id", "key", "title", "note"],
                    "field_schemas": {
                        field: {"dataType": "text", "displayName": field}
                        for field in ("key", "title", "note")
                    },
                    # The authoritative profile splits the import grants per
                    # operation instead of advertising their union twice.
                    "create_fields": ["key", "title"],
                    "update_fields": ["key", "note"],
                    "archive_field": None,
                    "date_updated_field": None,
                }
            if method == "import.preview":
                assert request["mode"] == "upsert"
                assert request["upsertKey"] == "key"
                rows_out = []
                for row in request["rows"]:
                    matched = row.get("key") in self.existing_keys
                    operation, grant_fields = (
                        ("update", self.update_fields)
                        if matched
                        else ("create", self.create_fields)
                    )
                    diagnostics = []
                    for field in row:
                        if field not in grant_fields:
                            diagnostics.append(
                                {
                                    "field": field,
                                    "code": "field.grant.denied",
                                    "message": f"field is not authorized by the {operation} grant",
                                }
                            )
                    rows_out.append({"values": row, "diagnostics": diagnostics})
                return {"rows": rows_out}
            if method == "import.mint":
                self.minted_rows = list(request["rows"])
                return {"token": "imp1.host-owned", "expiresAt": 9999999999}
            return await super().plugin_surface(principal, method, request, **kwargs)

    peer = LocalFilePeer()
    source = tmp_path / "upsert-grants.csv"
    # "note" is only writable through the update grant; "existing" matches a
    # record while "fresh" is a new insert row.
    source.write_text("key,note\nexisting,after\nfresh,new\n", encoding="utf-8")
    grant = peer.grants.issue(purpose="import_source", direction="read", path=str(source))
    peer.runs[grant.grant_id] = "host-run"
    client = HeterogeneousGrantClient()

    async def progress(*_args):
        pass

    plan = await preview_plugin_import(
        client,
        peer.files,
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "mode": "upsert",
            "upsertKey": "key",
            "pluginScope": scope(),
        },
        progress,
        lambda: False,
    )
    preview_call = next(
        request for _, method, request, _ in client.calls if method == "import.preview"
    )
    # The update-only column was routed through mapping into the authoritative preview.
    assert preview_call["rows"][0] == {"key": "existing", "note": "after"}
    assert preview_call["rows"][1] == {"key": "fresh", "note": "new"}
    # The matched update row keeps its granted update-only field clean.
    assert plan.rows[0].values == {"key": "existing", "note": "after"}
    assert plan.rows[0].diagnostics == []
    # The unmatched insert row is denied per row, not by dropping the column.
    assert [
        (diagnostic.code, diagnostic.column)
        for diagnostic in plan.rows[1].diagnostics
        if diagnostic.code == "field.grant.denied"
    ] == [("field.grant.denied", 2)]
    assert plan.summary.error_rows == 1
    # The minted plan still carries the routed values for the frozen Go plan path.
    assert [row["values"] for row in client.minted_rows] == [
        {"key": "existing", "note": "after"},
        {"key": "fresh", "note": "new"},
    ]


async def test_plugin_upsert_preview_requires_non_empty_upsert_key_before_authorization(tmp_path):
    """The plugin chain rejects an upsert without its approved key before any
    surface call, instead of previewing misleading updates the mint later rejects."""
    from backend.adapters.pocketbase.plugin_data_io import preview_plugin_import
    from backend.application.import_service import ImportFlowError

    peer = LocalFilePeer()
    source = tmp_path / "upsert-missing-key.csv"
    source.write_text("key,title\nexisting,keep\n", encoding="utf-8")
    grant = peer.grants.issue(purpose="import_source", direction="read", path=str(source))
    peer.runs[grant.grant_id] = "host-run"
    client = ScopedClient()

    async def progress(*_args):
        pass

    for raw in (
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "mode": "upsert",
            "pluginScope": scope(),
        },
        {
            "grantId": grant.grant_id,
            "collection": "orders",
            "schemaRevision": "schema_1",
            "mode": "upsert",
            "upsertKey": "",
            "pluginScope": scope(),
        },
    ):
        with pytest.raises(ImportFlowError) as error:
            await preview_plugin_import(client, peer.files, raw, progress, lambda: False)
        assert error.value.code == "plugin_capability_invalid"
    assert client.calls == [], "validation must precede any surface authorization"


async def test_bound_data_io_spool_rejects_oversize_before_commit_and_never_drops_run():
    calls = []

    async def call(action, params):
        calls.append((action, params))
        if action == "openWrite":
            return {"transferId": "t", "displayName": "output.csv"}
        assert action == "finishWrite"
        assert params["commit"] is False
        return {"committed": False, "bytes": 0}

    with pytest.raises(RuntimeError, match="byte budget"):
        async with HostFiles(call).bound_run("host-run", maximum_bytes=8).write("grant") as target:
            target.stream.write(b"123456789")
    assert calls[0][1] == {"grantId": "grant", "runId": "host-run"}
    assert all(action != "write" for action, _ in calls)


def test_plugin_workbook_expansion_is_bounded_before_openpyxl_parses():
    from zipfile import ZIP_DEFLATED, ZipFile

    from backend.application.import_service import ImportFlowError, SourceFile

    stream = io.BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("large-entry", b"x" * 1024)
    stream.seek(0)
    with pytest.raises(ImportFlowError) as error:
        SourceFile(stream, "source.xlsx", maximum_expanded_bytes=128).read_header_and_rows()
    assert error.value.code == "import_source_expansion_limit"


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize(
    ("logical_type", "fresh"),
    [
        ("json", True),
        ("json", False),
        ("formula", True),
        ("formula", False),
        ("lookup", True),
        ("lookup", False),
    ],
)
async def test_scoped_export_distinguishes_schema_computed_fields_from_json_values(
    tmp_path, fmt, logical_type, fresh
):
    import json

    from backend.adapters.pocketbase.plugin_data_io import export_plugin_data
    from backend.application.export_service import ExportError

    value = {
        "state": "ready" if logical_type != "json" else "draft",
        "value": "payload",
        "fresh": fresh,
    }

    class ShapeClient(ScopedClient):
        async def plugin_surface(self, principal, method, request, **kwargs):
            if method == "data.ioProfile":
                profile = await super().plugin_surface(principal, method, request, **kwargs)
                profile["field_schemas"] = {
                    "title": {"dataType": logical_type, "displayName": "title"}
                }
                return profile
            page = await super().plugin_surface(principal, method, request, **kwargs)
            page["items"] = [{"title": value}]
            page["totalRows"] = page["filteredRows"] = 1
            return page

    peer = LocalFilePeer()
    path = tmp_path / f"shape.{fmt}"
    grant = peer.grants.issue(purpose="export_target", direction="write", path=str(path))
    peer.runs[grant.grant_id] = "host-run"

    async def progress(*_args):
        pass

    operation = export_plugin_data(
        ShapeClient(),
        peer.files,
        {"grantId": grant.grant_id, "collection": "orders", "format": fmt, "pluginScope": scope()},
        progress,
        lambda: False,
    )
    if logical_type != "json" and not fresh:
        with pytest.raises(ExportError) as error:
            await operation
        assert error.value.code == "export_computed_cell_not_ready"
        assert not path.exists()
        return
    result = await operation
    assert result.rows_written == 1
    if fmt == "csv":
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
    else:
        book = load_workbook(path, read_only=True)
        try:
            rows = [list(row) for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
    cell = rows[1][0]
    if logical_type == "json":
        assert json.loads(cell) == value
    else:
        assert cell == "payload"
