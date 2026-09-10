from __future__ import annotations

import csv
import os
from collections.abc import Iterator, Mapping
from pathlib import Path
from typing import cast

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import JsonValue

from backend.__main__ import _configure_pocketbase_data_io
from backend.adapters.pocketbase.client import PocketBaseClient
from backend.adapters.pocketbase.transport import PocketBaseConfig, StdlibPocketBaseTransport
from backend.application.path_grant import SessionPathGrantStore
from backend.application.task_runtime import TaskRuntime
from backend.application.task_service import TaskService
from backend.contracts.product_rpc import JsonObject
from backend.contracts.task import (
    HostExportTargetParams,
    HostImportSourceParams,
    ResolveGrantParams,
)
from backend.rpc.dispatcher import RpcDispatcher
from backend.rpc.error_registry import ErrorDomain, register_application_errors
from tests.integration.packaged_sidecar_matrix import (
    CLAIM_ID,
    FENCE_EPOCH,
    SESSION_EPOCH,
    WORKSPACE_ID,
    Sidecar,
    _create_field,
    _create_table,
    _create_v2_workspace,
    _recommended_field_draft,
)
from tests.integration.test_data_io_interoperability_roundtrip import (
    source_sidecar_binary,  # noqa: F401 -- shared source-build fixture
)


@pytest.fixture(scope="module")
def grant_sidecar(
    request: pytest.FixtureRequest, tmp_path_factory: pytest.TempPathFactory
) -> Iterator[Sidecar]:
    sidecar = Sidecar(
        request.getfixturevalue("source_sidecar_binary"),
        _create_v2_workspace(tmp_path_factory.mktemp("path-grants") / "workspace"),
        workspace_identity={
            "VIBETABLE_WORKSPACE_ID": WORKSPACE_ID,
            "VIBETABLE_WORKSPACE_SESSION_EPOCH": str(SESSION_EPOCH),
            "VIBETABLE_WORKSPACE_FENCE_EPOCH": str(FENCE_EPOCH),
            "VIBETABLE_WORKSPACE_CLAIM_ID": CLAIM_ID,
        },
    )
    try:
        sidecar.start()
        yield sidecar
    finally:
        sidecar.stop()


class GrantDataIo:
    """Exercise the host registration and Data IO wire interface with real authority."""

    def __init__(self, sidecar: Sidecar, name: str) -> None:
        self.table = _create_table(sidecar, name, f"create-{name}")
        receipt = _create_field(
            sidecar,
            self.table,
            _recommended_field_draft(sidecar, self.table["tableId"], "Value", "text"),
            f"field-{name}",
        )
        self.field = receipt["definition"]["identity"]["physicalName"]
        config = PocketBaseConfig(
            base_url=f"http://{sidecar.address}", session_secret=sidecar.secret
        )
        self.client = PocketBaseClient(
            transport=StdlibPocketBaseTransport(config), session_secret=sidecar.secret
        )
        self.clock = [1000.0]
        tasks = TaskService(TaskRuntime(), SessionPathGrantStore(clock=lambda: self.clock[0]))
        self.rpc = RpcDispatcher()
        register_application_errors(ErrorDomain.PATH_GRANT)
        self.rpc.register(
            "path.registerImportSource", tasks.register_host_import_source, HostImportSourceParams
        )
        self.rpc.register(
            "path.registerExportTarget", tasks.register_host_export_target, HostExportTargetParams
        )
        self.rpc.register("path.resolveGrant", tasks.resolve_grant, ResolveGrantParams)
        _configure_pocketbase_data_io(self.rpc, client=self.client, task_service=tasks)

    async def call(self, method: str, params: JsonObject) -> JsonObject:
        response = await self.rpc.dispatch(
            {"jsonrpc": "2.0", "id": method, "method": method, "params": params}
        )
        assert response is not None
        assert response["id"] == method
        return cast(JsonObject, response)

    async def result(self, method: str, params: JsonObject) -> JsonObject:
        response = await self.call(method, params)
        assert "error" not in response, response
        return cast(JsonObject, response["result"])

    async def register(self, path: Path, *, export: bool = False) -> str:
        descriptor = await self.result(
            "path.registerExportTarget" if export else "path.registerImportSource",
            {"path": str(path)},
        )
        assert "path" not in descriptor
        assert descriptor["displayName"] == path.name
        return str(descriptor["grantId"])

    async def preview(self, grant: str) -> JsonObject:
        return await self.result(
            "data.previewImport",
            {
                "grantId": grant,
                "collection": self.table["tableId"],
                "schemaRevision": self.table["schemaRevision"],
            },
        )

    def apply_params(self, grant: str, plan: JsonObject, prefix: str) -> JsonObject:
        return {
            "grantId": grant,
            "collection": self.table["tableId"],
            "token": cast(JsonObject, plan["token"])["token"],
            "idempotencyPrefix": prefix,
        }

    async def authority(self) -> object:
        page = await self.client.query_page(table_id=self.table["tableId"], query={"limit": 100})
        return (
            page.rows,
            page.total_rows,
            page.filtered_rows,
            page.snapshot["schemaRevision"],
            page.snapshot["dataRevision"],
        )


@pytest.mark.integration
@pytest.mark.asyncio
@pytest.mark.parametrize("invalidity", ["expired", "different-grant", "consumed"])
async def test_apply_rejects_invalid_grant_before_authority_mutation(
    grant_sidecar: Sidecar,
    tmp_path: Path,
    invalidity: str,
) -> None:
    io = GrantDataIo(grant_sidecar, "grant_" + invalidity.replace("-", "_"))
    source = tmp_path / "source.csv"
    source.write_text(f"{io.field}\n授权值\n", encoding="utf-8-sig")
    grant = await io.register(source)
    plan = await io.preview(grant)
    submitted_grant = grant
    if invalidity == "expired":
        io.clock[0] += 300.0
    elif invalidity == "different-grant":
        submitted_grant = await io.register(source)
    else:
        other_plan = await io.preview(grant)
        await io.result("data.applyImport", io.apply_params(grant, other_plan, "consume-grant"))
    before = await io.authority()
    response = await io.call(
        "data.applyImport", io.apply_params(submitted_grant, plan, "reject-" + invalidity)
    )
    # A failure response alone is insufficient: the old implementation committed
    # first and only discovered an expired/consumed grant afterwards.
    assert await io.authority() == before
    assert "error" in response, response
    error = cast(JsonObject, response["error"])
    assert error["code"] == (-32060 if invalidity == "different-grant" else -32050)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_admitted_import_commits_once_when_grant_expires_after_authority_receipt(
    grant_sidecar: Sidecar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    io = GrantDataIo(grant_sidecar, "grant_expires_after_commit")
    source = tmp_path / "expires-after-commit.csv"
    source.write_text(f"{io.field}\n只写入一次\n", encoding="utf-8-sig")
    grant = await io.register(source)
    plan = await io.preview(grant)
    another_plan = await io.preview(grant)
    real_apply = io.client.apply_mutation
    mutation_calls = 0

    async def expire_after_receipt(request: Mapping[str, JsonValue]) -> dict[str, JsonValue]:
        nonlocal mutation_calls
        receipt = await real_apply(request)
        mutation_calls += 1
        # Cross the TTL only after the real sidecar has answered. The response
        # remains untouched; this varies time at the existing client seam.
        io.clock[0] += 301.0
        return receipt

    monkeypatch.setattr(io.client, "apply_mutation", expire_after_receipt)
    params = io.apply_params(grant, plan, "expiry-after-commit")
    applied = await io.result("data.applyImport", params)
    assert applied["createdCount"] == 1
    assert applied["failedRows"] == []
    page = await io.client.query_page(table_id=io.table["tableId"], query={"limit": 100})
    assert len(page.rows) == 1
    assert page.rows[0][io.field] == "只写入一次"
    before = await io.authority()
    descriptor = await io.call("path.resolveGrant", {"grantId": grant})
    assert cast(JsonObject, descriptor["error"])["code"] == -32050
    replay = await io.call("data.applyImport", params)
    assert cast(JsonObject, replay["error"])["code"] == -32060
    other = await io.call(
        "data.applyImport", io.apply_params(grant, another_plan, "expiry-second-plan")
    )
    assert cast(JsonObject, other["error"])["code"] == -32050
    assert mutation_calls == 1
    assert await io.authority() == before


@pytest.mark.integration
@pytest.mark.asyncio
@pytest.mark.skipif(os.name != "nt", reason="Windows extended-length filesystem qualification")
@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
async def test_windows_long_path_host_grant_import_export_and_replay(
    grant_sidecar: Sidecar,
    tmp_path: Path,
    export_format: str,
) -> None:
    io = GrantDataIo(grant_sidecar, "long_path_" + export_format)
    directory = tmp_path
    while len(str(directory)) < 280:
        directory /= "中文-Cafe\u0301-" + "a" * 35
    directory.mkdir(parents=True)
    source = directory / f"源-👩‍💻.{export_format}"
    expected = "原样 Cafe\u0301 👩‍💻"
    if export_format == "csv":
        source.write_text(f"{io.field}\n{expected}\n", encoding="utf-8-sig")
    else:
        workbook = Workbook()
        assert workbook.active is not None
        workbook.active.append([io.field])
        workbook.active.append([expected])
        workbook.save(source)
        workbook.close()
    grant = await io.register(source)
    plan = await io.preview(grant)
    params = io.apply_params(grant, plan, "long-path-" + export_format)
    applied = await io.result("data.applyImport", params)
    assert applied["createdCount"] == 1
    before = await io.authority()
    descriptor = await io.call("path.resolveGrant", {"grantId": grant})
    assert cast(JsonObject, descriptor["error"])["code"] == -32050
    preview = await io.call(
        "data.previewImport",
        {
            "grantId": grant,
            "collection": io.table["tableId"],
            "schemaRevision": io.table["schemaRevision"],
        },
    )
    assert cast(JsonObject, preview["error"])["code"] == -32050
    replay = await io.call("data.applyImport", params)
    assert cast(JsonObject, replay["error"])["code"] == -32060
    target = directory / f"导出-👩‍💻.{export_format}"
    target_grant = await io.register(target, export=True)
    exported = await io.result(
        "data.export",
        {
            "grantId": target_grant,
            "collection": io.table["tableId"],
            "query": {},
            "format": export_format,
        },
    )
    assert exported["rowsWritten"] == 1
    if export_format == "csv":
        with target.open(encoding="utf-8-sig", newline="") as stream:
            assert next(csv.DictReader(stream))[io.field] == expected
    else:
        workbook = load_workbook(target, read_only=True)
        try:
            assert workbook.active is not None
            rows = workbook.active.iter_rows(values_only=True)
            headers = list(next(rows))
            assert next(rows)[headers.index(io.field)] == expected
        finally:
            workbook.close()
    assert await io.authority() == before
    assert sorted(path.name for path in directory.iterdir()) == sorted([source.name, target.name])
