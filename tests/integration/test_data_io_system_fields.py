from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook

from backend.adapters.pocketbase.client import PocketBaseClient
from backend.adapters.pocketbase.data_io import ProductDataIoRuntime
from backend.adapters.pocketbase.transport import PocketBaseConfig, StdlibPocketBaseTransport
from backend.application.revisioned_metadata_port import JsonObject
from backend.application.task_service import build_task_service
from backend.contracts.data_io import (
    ApplyImportParams,
    ExportParams,
    ImportColumnMapping,
    PreviewImportParams,
)
from backend.contracts.task import HostExportTargetParams, HostImportSourceParams
from tests.integration.packaged_sidecar_matrix import (
    CLAIM_ID,
    FENCE_EPOCH,
    SESSION_EPOCH,
    WORKSPACE_ID,
    Sidecar,
    _apply,
    _create_field,
    _create_table,
    _create_v2_workspace,
    _recommended_field_draft,
)
from tests.integration.test_data_io_interoperability_roundtrip import (
    source_sidecar_binary as source_sidecar_binary,
)

REPO_ROOT = Path(__file__).resolve().parents[2]
CORPUS_PATH = REPO_ROOT / "tests" / "fixtures" / "data-io" / "a5-system-field-corpus.json"


def _load_corpus() -> dict[str, Any]:
    corpus = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    assert corpus["$schema"] == "vibetable.a5-system-field-corpus.v1"
    assert {item["format"] for item in corpus["inputs"]} == {"csv", "xlsx"}
    return cast(dict[str, Any], corpus)


def _auto_date_draft(
    sidecar: Sidecar, table_id: str, display_name: str, role: str
) -> dict[str, Any]:
    draft = _recommended_field_draft(sidecar, table_id, display_name, "autoDate")
    draft["autoDate"] = {"role": role}
    return draft


def _write_source(path: Path, source_format: str, headers: list[str], values: list[str]) -> None:
    if source_format == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            writer.writerow(headers)
            writer.writerow(values)
        return
    workbook = Workbook()
    try:
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.append(headers)
        worksheet.append(values)
        workbook.save(path)
    finally:
        workbook.close()


def _read_export(path: Path, export_format: str) -> dict[str, object]:
    if export_format == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return dict(next(csv.DictReader(stream)))
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        assert worksheet is not None
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows))
        values = list(next(rows))
        return dict(zip(headers, values, strict=True))
    finally:
        workbook.close()


@pytest.mark.integration
@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("source_format", "use_explicit_mapping"),
    [("csv", False), ("xlsx", True)],
    ids=("csv-auto-map", "xlsx-explicit-map"),
)
async def test_system_fields_are_excluded_from_import_and_remain_authority_owned(
    tmp_path: Path,
    source_sidecar_binary: Path,
    source_format: str,
    use_explicit_mapping: bool,
) -> None:
    corpus = _load_corpus()
    input_case = next(item for item in corpus["inputs"] if item["format"] == source_format)
    forged = cast(dict[str, str], corpus["forged"])
    sidecar = Sidecar(
        source_sidecar_binary,
        _create_v2_workspace(tmp_path / "workspace"),
        workspace_identity={
            "VIBETABLE_WORKSPACE_ID": WORKSPACE_ID,
            "VIBETABLE_WORKSPACE_SESSION_EPOCH": str(SESSION_EPOCH),
            "VIBETABLE_WORKSPACE_FENCE_EPOCH": str(FENCE_EPOCH),
            "VIBETABLE_WORKSPACE_CLAIM_ID": CLAIM_ID,
        },
    )
    try:
        sidecar.start()
        table = _create_table(sidecar, "a5_system_fields", "a5-create-system-table")
        text = _create_field(
            sidecar,
            table,
            _recommended_field_draft(sidecar, table["tableId"], "Text", "text"),
            "a5-create-system-text",
        )["definition"]
        created = _create_field(
            sidecar,
            table,
            _auto_date_draft(sidecar, table["tableId"], "Created", "createdAt"),
            "a5-create-system-created",
        )["definition"]
        updated = _create_field(
            sidecar,
            table,
            _auto_date_draft(sidecar, table["tableId"], "Updated", "updatedAt"),
            "a5-create-system-updated",
        )["definition"]
        text_field = text["identity"]["physicalName"]
        created_field = created["identity"]["physicalName"]
        updated_field = updated["identity"]["physicalName"]
        headers = ["id", created_field, updated_field, text_field]
        source = tmp_path / f"a5-system-fields.{source_format}"
        _write_source(
            source,
            source_format,
            headers,
            [forged["id"], forged["createdAt"], forged["updatedAt"], input_case["text"]],
        )

        config = PocketBaseConfig(
            base_url=f"http://{sidecar.address}", session_secret=sidecar.secret
        )
        client = PocketBaseClient(
            transport=StdlibPocketBaseTransport(config), session_secret=sidecar.secret
        )
        tasks = build_task_service()
        runtime = ProductDataIoRuntime(client=client, task_service=tasks)
        grant = await tasks.register_host_import_source(
            HostImportSourceParams(
                path=str(source.resolve()),
                size_bytes=source.stat().st_size,
                mime_type="text/csv"
                if source_format == "csv"
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        )
        mappings = []
        if use_explicit_mapping:
            mappings = [
                ImportColumnMapping(source_column=header, target_field=header) for header in headers
            ]
        plan = await runtime.preview_import(
            PreviewImportParams(
                grant_id=grant.grant_id,
                collection=table["tableId"],
                schema_revision=table["schemaRevision"],
                column_mapping=mappings,
            )
        )
        assert plan.summary.total_rows == plan.summary.valid_rows == 1
        assert plan.summary.error_count == 0
        assert plan.rows[0].values == {text_field: input_case["text"]}
        assert plan.unmatched_columns == ["id", created_field, updated_field]

        query: JsonObject = {"filters": [], "sorts": [], "offset": 0, "limit": 100}
        applied = await runtime.apply_import(
            ApplyImportParams(
                grant_id=grant.grant_id,
                collection=table["tableId"],
                token=plan.token.token,
                idempotency_prefix=f"a5-system-{source_format}",
            )
        )
        assert applied.created_count == 1
        assert applied.failed_rows == []
        authority = await client.query_page(table_id=table["tableId"], query=query)
        assert len(authority.rows) == 1
        authority_row = authority.rows[0]
        assert authority_row["id"] != forged["id"]
        assert authority_row[text_field] == input_case["text"]
        for field, forged_value in (
            (created_field, forged["createdAt"]),
            (updated_field, forged["updatedAt"]),
        ):
            assert isinstance(authority_row[field], str)
            assert authority_row[field] != forged_value

        for export_format in ("csv", "xlsx"):
            target = tmp_path / f"a5-system-export.{export_format}"
            export_grant = await tasks.register_host_export_target(
                HostExportTargetParams(path=str(target.resolve()))
            )
            result = await runtime.export(
                ExportParams(
                    grant_id=export_grant.grant_id,
                    collection=table["tableId"],
                    query=query,
                    format=export_format,
                )
            )
            assert result.rows_written == 1
            exported = _read_export(target, export_format)
            assert exported[created_field] == authority_row[created_field]
            assert exported[updated_field] == authority_row[updated_field]

        after_export = await client.query_page(table_id=table["tableId"], query=query)
        assert after_export.rows == authority.rows
        for revision_key in ("table", "schemaRevision", "dataRevision"):
            assert after_export.snapshot[revision_key] == authority.snapshot[revision_key]
        rejected = _apply(
            sidecar,
            table["tableId"],
            table["schemaRevision"],
            f"a5-reject-system-{source_format}",
            [
                {
                    "kind": "update",
                    "recordId": authority_row["id"],
                    "values": {created_field: forged["createdAt"]},
                }
            ],
            expected=422,
        )
        assert rejected["code"] == "mutation.field.read_only"
        after_rejection = await client.query_page(table_id=table["tableId"], query=query)
        assert after_rejection.rows == authority.rows
        for revision_key in ("table", "schemaRevision", "dataRevision"):
            assert after_rejection.snapshot[revision_key] == authority.snapshot[revision_key]
    finally:
        sidecar.stop()
