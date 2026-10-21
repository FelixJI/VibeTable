"""Native fixture and fail-closed row/type contracts for the real S35 exports."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.application.import_service import SourceFile
from tests.e2e.data_io_workbook import verify_export, write_native

COLUMNS = ["note", "stamp"]
ROWS = [["=文本", "2026-08-29 14:05:06.123Z"], ["普通文本", "2026-08-28 16:00:00.000Z"]]

# Use a macroEnabled workbook content type rather than renaming an .xlsx file.
# This fixture contains no VBA project and covers reading values only.
MACRO_MAIN_CONTENT_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
PLAIN_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


def write_macro_enabled_workbook(path: Path, note: str) -> None:
    """Create a macroEnabled workbook without a VBA project for reading tests."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["day", "stamp", "note"])
    sheet.append([date(2026, 8, 29), datetime(2026, 8, 29, 14, 5, 6, 123000), note])
    workbook.save(path)
    workbook.close()
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    types = parts["[Content_Types].xml"].decode("utf-8")
    assert PLAIN_MAIN_CONTENT_TYPE in types
    types = types.replace(PLAIN_MAIN_CONTENT_TYPE, MACRO_MAIN_CONTENT_TYPE)
    parts["[Content_Types].xml"] = types.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)


def test_native_fixture_preserves_dates_milliseconds_and_plain_formula_text(tmp_path: Path) -> None:
    target = tmp_path / "native.xlsx"
    write_native(
        target,
        ["day", "stamp", "note"],
        "2026-08-29",
        "2026-08-29 14:05:06.123",
        ["=文本", "普通文本"],
    )
    workbook = load_workbook(target, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        assert list(sheet.values) == [
            ("day", "stamp", "note"),
            (datetime(2026, 8, 29), datetime(2026, 8, 29, 14, 5, 6, 123000), "=文本"),
            (datetime(2026, 8, 29), datetime(2026, 8, 29, 14, 5, 6, 123000), "普通文本"),
        ]
        for row in sheet.iter_rows(min_row=2):
            assert [cell.data_type for cell in row] == ["d", "d", "s"]
    finally:
        workbook.close()


def write_export(target: Path, rows: list[list[str]], columns: list[str] = COLUMNS) -> None:
    if target.suffix == ".csv":
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows([columns, *rows])
    else:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        for values in [columns, *rows]:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
        workbook.save(target)
        workbook.close()


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_export_verification_uses_columns_and_unordered_exact_rows(
    tmp_path: Path, suffix: str
) -> None:
    target = tmp_path / f"export{suffix}"
    write_export(target, [list(reversed(row)) for row in reversed(ROWS)], list(reversed(COLUMNS)))
    assert verify_export(target, COLUMNS, ROWS)["rows"] == 2


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("corruption", ["swap_dates", "duplicate_note", "missing", "extra"])
def test_export_verification_rejects_row_corruption(
    tmp_path: Path, suffix: str, corruption: str
) -> None:
    rows = [row.copy() for row in ROWS]
    if corruption == "swap_dates":
        rows[0][1], rows[1][1] = rows[1][1], rows[0][1]
    elif corruption == "duplicate_note":
        rows[1][0] = rows[0][0]
    elif corruption == "missing":
        rows.pop()
    else:
        rows.append(rows[0].copy())
    target = tmp_path / f"corrupt{suffix}"
    write_export(target, rows)
    with pytest.raises(AssertionError, match="row multiset"):
        verify_export(target, COLUMNS, ROWS)


@pytest.mark.parametrize("value", ["=1+1", 123, datetime(2026, 8, 29)])
def test_xlsx_verification_rejects_non_string_cells(tmp_path: Path, value: object) -> None:
    target = tmp_path / "typed.xlsx"
    write_export(target, ROWS)
    workbook = load_workbook(target)
    workbook.worksheets[0].cell(2, 2).value = value
    workbook.save(target)
    workbook.close()
    with pytest.raises(AssertionError, match="must be strings"):
        verify_export(target, COLUMNS, ROWS)


def test_macro_enabled_xlsm_reads_through_real_import_service(tmp_path: Path) -> None:
    """Read values from a macroEnabled workbook without VBA through the importer."""
    path = tmp_path / "数据-源-宏.xlsm"
    write_macro_enabled_workbook(path, "宏工作簿-👩🏽‍💻")

    # The container really is macroEnabled, not a renamed plain workbook.
    with zipfile.ZipFile(path) as archive:
        types = archive.read("[Content_Types].xml").decode("utf-8")
        assert MACRO_MAIN_CONTENT_TYPE in types
        assert PLAIN_MAIN_CONTENT_TYPE not in types
        assert "xl/vbaProject.bin" not in archive.namelist()

    header, rows, _ = SourceFile(io.BytesIO(path.read_bytes()), path.name).read_header_and_rows()
    assert header == ["day", "stamp", "note"]
    assert rows == [["2026-08-29", "2026-08-29 14:05:06.123000", "宏工作簿-👩🏽‍💻"]]

    # A plain .xlsx sibling with the same cells must read identically: the
    # macroEnabled container takes the same reading path with no value drift.
    sibling = tmp_path / "数据-源-plain.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(header)
    sheet.append([date(2026, 8, 29), datetime(2026, 8, 29, 14, 5, 6, 123000), "宏工作簿-👩🏽‍💻"])
    workbook.save(sibling)
    workbook.close()
    plain_header, plain_rows, _ = SourceFile(
        io.BytesIO(sibling.read_bytes()), sibling.name
    ).read_header_and_rows()
    assert (plain_header, plain_rows) == (header, rows)
