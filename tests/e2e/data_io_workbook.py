"""S35 fixture producer and independent tabular export verifier; no product calls."""

from __future__ import annotations

import csv
import json
import locale
import platform
import sys
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import et_xmlfile
import openpyxl
from openpyxl import Workbook, load_workbook


def write_native(
    path: Path, columns: list[str], day: str, stamp: str, notes: list[str]
) -> dict[str, object]:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(columns)
    for note in notes:
        sheet.append([date.fromisoformat(day), datetime.fromisoformat(stamp), note])
        sheet.cell(sheet.max_row, 3).data_type = "s"
    workbook.save(path)
    workbook.close()
    return {
        "python": platform.python_version(),
        "openpyxl": openpyxl.__version__,
        "etXmlfile": et_xmlfile.__version__,
        "locale": locale.setlocale(locale.LC_ALL),
        "encoding": locale.getencoding(),
        "timezoneNames": list(time.tzname),
        "utcOffsetSeconds": -time.timezone,
        "xlsxEpoch": "1900",
        "csvEncoding": "utf-8-sig",
        "source": str(path),
    }


def verify_export(path: Path, columns: list[str], expected: list[list[str]]) -> dict[str, object]:
    if path.suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
    elif path.suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            assert len(workbook.worksheets) == 1, "Expected exactly one exported worksheet"
            rows = []
            for cells in workbook.worksheets[0].iter_rows():
                # All populated cells in these text/date-wire fixtures are strings.
                # data_only=False ensures formulas cannot masquerade as cached text.
                assert all(cell.data_type == "s" for cell in cells if cell.value is not None), (
                    "Exported cells must be strings, not formulas, numbers or native dates"
                )
                rows.append([cell.value for cell in cells])
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported fixture export format")
    assert rows, "Missing header"
    header, *data = rows
    assert len(header) == len(set(header)), "Duplicate export header"
    assert all(column in header for column in columns), "Missing export column"
    assert all(len(row) == len(header) for row in data), "Ragged exported rows"
    indexes = [header.index(column) for column in columns]
    actual = [tuple(row[index] for index in indexes) for row in data]
    assert Counter(actual) == Counter(map(tuple, expected)), "Exported row multiset differs"
    return {"rows": len(actual), "columns": columns, "format": path.suffix[1:]}


def main() -> None:
    action, filename, payload_text = sys.argv[1:]
    payload = json.loads(payload_text)
    path = Path(filename)
    if action == "native":
        result = write_native(
            path, payload["columns"], payload["date"], payload["stamp"], payload["notes"]
        )
    elif action == "verify":
        result = verify_export(path, payload["columns"], payload["rows"])
    else:
        raise ValueError("Unsupported workbook fixture action")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
