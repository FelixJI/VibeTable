"""The export oracle must reject wrong values/order and executable formulas."""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from tests.e2e.data_io_workbook import verify_export


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_calculation_export_preserves_numeric_values_and_order(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"chain{suffix}"
    columns = ["contract", "amount", "text"]
    expected: list[list[str | int | float]] = [["甲", 22, "=1+1"], ["乙", 1, "literal"]]
    if suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerows([columns, *expected])
    else:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(columns)
        for row in expected:
            sheet.append(row)
            sheet.cell(sheet.max_row, 3).data_type = "s"
        workbook.save(path)
        workbook.close()
    assert verify_export(path, columns, expected, string_cells_only=False)["rows"] == 2
    for wrong in [list(reversed(expected)), [["甲", 21, "=1+1"], expected[1]], [expected[0]]]:
        with pytest.raises(AssertionError):
            verify_export(path, columns, wrong, string_cells_only=False)
    if suffix == ".xlsx":
        with pytest.raises(AssertionError):
            verify_export(path, columns, expected)  # S35 remains string-only.


def test_calculation_export_rejects_executable_xlsx_formula(tmp_path: Path) -> None:
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["text"])
    sheet.append(["=1+1"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(AssertionError):
        verify_export(path, ["text"], [["=1+1"]], string_cells_only=False)
